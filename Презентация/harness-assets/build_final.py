# -*- coding: utf-8 -*-
"""Сборка деки на каркасе рабочего шаблона владельца: клонирование слайдов + заполнение."""
import zipfile, re, shutil
from pptx import Presentation
from pptx.util import Emu

SRC = "D:/IA/tmp-harness/template.pptx"
BASE = "D:/IA/tmp-harness/base12.pptx"
OUT = "D:/IA/ace/Презентация/harness-assets/harness-deck-v2.pptx"
ASSET = "D:/IA/ace/Презентация/harness-assets/"

# порядок слайдов: (источник, роль)
ORDER = [1, 3, 4, 3, 3, 3, 3, 4, 3, 3, 4, 5]

# ── 1. Пересборка пакета ─────────────────────────────────
zin = zipfile.ZipFile(SRC)
names = zin.namelist()
drop_prefixes = ("ppt/slides/", "ppt/notesSlides/")
zout = zipfile.ZipFile(BASE, "w", zipfile.ZIP_DEFLATED)
for n in names:
    if n.startswith(drop_prefixes) or n == "[Content_Types].xml" or n in ("ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"):
        continue
    zout.writestr(n, zin.read(n))

# слайды
for i, src in enumerate(ORDER, 1):
    zout.writestr(f"ppt/slides/slide{i}.xml", zin.read(f"ppt/slides/slide{src}.xml"))
    zout.writestr(f"ppt/slides/_rels/slide{i}.xml.rels", zin.read(f"ppt/slides/_rels/slide{src}.xml.rels"))

# presentation rels: rId1 master, rId2 theme, rId3-5 props, слайды rId11+
rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="slideMasters/slideMaster1.xml"/>',
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="theme/theme1.xml"/>']
for i in range(12):
    rels.append(f'<Relationship Id="rId{11+i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i+1}.xml"/>')
rels.append('<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/presProps" Target="presProps.xml"/>')
rels.append('<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/viewProps" Target="viewProps.xml"/>')
rels.append('<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/tableStyles" Target="tableStyles.xml"/>')
rels.append("</Relationships>")
zout.writestr("ppt/_rels/presentation.xml.rels", "".join(rels))

# presentation.xml: заменить sldIdLst
pres = zin.read("ppt/presentation.xml").decode("utf-8")
sldids = "".join(f'<p:sldId id="{256+i}" r:id="rId{11+i}"/>' for i in range(12))
pres = re.sub(r"<p:sldIdLst>.*?</p:sldIdLst>", f"<p:sldIdLst>{sldids}</p:sldIdLst>", pres)
zout.writestr("ppt/presentation.xml", pres)

# content types
ct = zin.read("[Content_Types].xml").decode("utf-8")
ct = re.sub(r"<Override[^>]*ppt/slides/slide\d+\.xml[^>]*/>", "", ct)
slide_ct = '<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
overrides = "".join(slide_ct.format(i=i) for i in range(1, 13))
ct = ct.replace("</Types>", overrides + "</Types>")
zout.writestr("[Content_Types].xml", ct)
zout.close()
print("base12 written")

# ── 2. Заполнение контента ───────────────────────────────
def set_text(shape, text):
    tf = shape.text_frame
    paras = tf.paragraphs
    p0 = paras[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.add_run().text = text
    for p in paras[1:]:
        p._p.getparent().remove(p._p)

def set_lines(shape, lines):
    tf = shape.text_frame
    for i, p in enumerate(tf.paragraphs):
        if i < len(lines) and p.runs:
            p.runs[0].text = lines[i]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        elif i < len(lines):
            p.add_run().text = lines[i]

def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None

def swap_screen(slide, img):
    scr = find(slide, "screen")
    l, t, w, h = scr.left, scr.top, scr.width, scr.height
    scr._element.getparent().remove(scr._element)
    slide.shapes.add_picture(ASSET + img, l, t, w, h)

def drop_shape(slide, name):
    sh = find(slide, name)
    if sh is not None:
        sh._element.getparent().remove(sh._element)

prs = Presentation(BASE)
S = list(prs.slides)

def fill_step(slide, tag, title, bullets, img=None, caption=None, pagenum=None):
    set_text(find(slide, "tag"), tag)
    set_text(find(slide, "title"), title)
    for i, b in enumerate(bullets, 1):
        set_text(find(slide, f"bullet-{i}"), b)
    if img:
        swap_screen(slide, img)
    if caption:
        set_text(find(slide, "screen-caption"), caption)
    if pagenum:
        set_text(find(slide, "page-num"), pagenum)

def fill_kpi(slide, tag, title, kpis, subs, note, pagenum):
    set_text(find(slide, "tag"), tag)
    set_text(find(slide, "title"), title)
    for i, (v, l) in enumerate(kpis, 1):
        set_text(find(slide, f"kpi-val-{i}"), v)
        set_text(find(slide, f"kpi-lbl-{i}"), l)
    for i, (v, l) in enumerate(subs, 1):
        set_text(find(slide, f"sub-val-{i}"), v)
        set_text(find(slide, f"sub-lbl-{i}"), l)
    set_text(find(slide, "note-text"), note)
    set_text(find(slide, "page-num"), pagenum)

# S1 титул
s = S[0]
set_text(s, "subtitle") if False else set_text(find(s, "subtitle"),
    "Один инженер и AI-агент: от формул и нормативки — до установщика, отчётов и облака проверок")
set_text(find(s, "meta"), "ИИ-конвейер · v1.2.0 · сентябрь 2026")

# S2 продукт
fill_step(S[1], "ПРОДУКТ · ЧТО ДЕЛАЕТ ПРИЛОЖЕНИЕ", "Полный цикл: от климата до записки",
    ["Расчёт снеготаяния — тепловая нагрузка по климату города",
     "Гидравлика коллекторов — диаметры, скорости, балансировка",
     "Подбор оборудования РЕХАУ — по каталогу, с ограничениями",
     "Пояснительная записка в PDF — одним кликом, для экспертизы"],
    "render-hydraulics.png", "Гидравлический расчёт: контуры, коллекторы, дросселирование", "02")

# S3 технологии
fill_kpi(S[2], "СТЕК · ТЕХНОЛОГИИ ВНУТРИ", "Стек вендора + ИИ-инструментарий",
    [(".NET 8", "язык и рантайм от Microsoft"), ("WPF", "настольный UI-фреймворк Windows"),
     ("2 237", "автотестов NUnit на каждый расчёт"), ("MIT", "свободная лицензия PDF-движка")],
    [("GitHub Actions", "CI: проверка каждой сборки"), ("ZCode · GLM", "AI-агент: план, код, ревью"),
     ("Inno Setup", "Windows-установщик"), ("Брендбук", "токены дизайн-системы")],
    "Ни одного компонента с лицензионным риском: каждый выбор зафиксирован в журнале решений проекта.", "03")

# S4 конвейер
fill_step(S[3], "КОНВЕЙЕР · КАК ДЕЛАЕТСЯ ФИЧА", "Шесть шагов — и ни один без гейта",
    ["План и ревью — второй ИИ ищет дефекты до написания кода",
     "Show-me — владелец утверждает картинку будущего",
     "Реализация с тестами — 2 237 проверок на каждом шаге",
     "Независимое ревью диффа — затем приёмка человеком"],
    "conv-chain.png", "11 решений до кода · 3 дефекта нашло ревью · +22 теста · 0 регрессий", "04")

# S5 правила
fill_step(S[4], "ХАРНЕСС · ПРАВИЛА ДЛЯ ИИ", "Устав проекта обеспечен машиной",
    ["AGENTS.md — правила архитектуры и процесса в репозитории",
     "ADR-журнал — 14 решений с причинами и последствиями",
     "Машинные проверки — инварианты превращены в тесты",
     "Списки писателей меняются только через ADR-запись"],
    "terminal.png", "Нарушение правила падает тестом — автоматически", "05")

# S6 память
fill_step(S[5], "ХАРНЕСС · ПАМЯТЬ ПРОЕКТА", "ИИ не начинает с нуля",
    ["Журнал решений: 14 ADR — почему система устроена так",
     "Журнал уроков: 21 запись — ошибку не повторить «по незнанию»",
     "Коммиты-досье: история каждой фичи читается как отчёт",
     "Новая сессия ИИ читает историю и продолжает работу"],
    "adr-tape.png", "Реальные записи: ADR-009, ADR-012, ADR-014", "06")

# S7 телефон
fill_step(S[6], "МОБИЛЬНОСТЬ · ПУЛЬТ В КАРМАНЕ", "Разработка управляется с телефона",
    ["ZCode Remote Control — сессия агента видна в реальном времени",
     "Планы и визуализации открываются на экране телефона",
     "Приёмка фазы — коротким сообщением, где бы вы ни были",
     "На скрине — живые задачи этого проекта"],
    "phone-anim.gif", "Живые задачи проекта: расчёты, редизайны, дедлоки", "07")

# S8 надёжность
fill_kpi(S[7], "КАЧЕСТВО · НАДЁЖНОСТЬ В ЦИФРАХ", "Каждая цифра подкреплена тестом",
    [("2 237", "автотестов: код, расчёты, правила"), ("11", "фаз перестройки без потери данных"),
     ("10", "шагов отмены любых правок"), ("0", "регрессий в кейсе undo/redo")],
    [("всегда", "старые файлы открываются в новых"), ("ru-RU", "каноническая локаль всех чисел"),
     ("UiSmoke", "интерфейс проверяется тестами"), ("Ревью", "перед серьёзными изменениями")],
    "Источник цифр — репозиторий проекта: тесты, журнал ADR, история изменений (сентябрь 2026).", "08")

# S9 CI
fill_step(S[8], "АВТОМАТИЗАЦИЯ · ИНФРАСТРУКТУРА КАЧЕСТВА", "GitHub проверяет каждую сборку",
    ["CI в облаке: сборка плюс 2 237 тестов на каждый коммит",
     "Зелёная галочка — к приёмке, красный крест — обратно агенту",
     "Реестр «правило → проверка» — правило не может протухнуть",
     "Ревью-чеки — у каждого ревью перечитываемый вердикт"],
    "ci-flow.png", "Контур: изменение → GitHub → CI → приёмка", "09")

# S10 рост
fill_step(S[9], "РАЗВИТИЕ · КУДА РАСТУ", "Следующие модули того же конвейера",
    ["Telegram-приёмка — аппрувы фаз прямо в мессенджере",
     "RAG: агент сам ищет в брендбуке и нормативке — с цитатой",
     "MCP: GitHub, сборки, задачи — как инструменты агента",
     "Автономные сценарии: релиз, разбор сбоев, метрики"],
    "growth.png", "Все технологии — отраслевые стандарты", "10")

# S11 масштаб
fill_kpi(S[10], "МАСШТАБИРОВАНИЕ · ПЕРЕНОС НА КОМАНДУ", "Тот же контур работает у отдела",
    [("01", "Правила — письменно, в репозитории"), ("02", "Проверки — машинные, тестами"),
     ("03", "Сборки — в облаке, автоматически"), ("04", "Ревью — ИИ проверяет ИИ")],
    [("День", "старт — обычный текстовый файл"), ("₽0", "GitHub и CI — бесплатный тариф"),
     ("1 → N", "подходит и одному, и отделу"), ("Человек", "цель ставит и принимает результат")],
    "Четыре шага внедряются по очереди — каждый уже опробован в этом проекте.", "11")

# S12 финал — без декоративного паттерна; фирменные «квадраты» (стр. 41/57) как уголок
s = S[11]
drop_shape(s, "pattern-right")
s.shapes.add_picture(ASSET + "squares-corner.png", Emu(9906000), Emu(4876800), height=Emu(1828800))
set_lines(find(s, "contacts"),
    ["Готов к пилоту — проверим на реальных задачах проектировщиков",
     "Калькулятор снеготаяния · v1.2.0 · сентябрь 2026"])
set_text(find(s, "page-num"), "12")

prs.core_properties.author = "РЕХАУ"
prs.core_properties.title = "Калькулятор снеготаяния — ИИ-конвейер разработки"
prs.save(OUT)

# пост-обработка: фирменный шрифт Inter вместо шаблонного MiSans (брендбук, стр. 19)
import zipfile
tmp = OUT + ".tmp"
zin = zipfile.ZipFile(OUT)
zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
for n in zin.namelist():
    data = zin.read(n)
    if n.endswith(".xml"):
        data = data.replace(b'typeface="MiSans"', b'typeface="Inter"')
    zout.writestr(n, data)
zout.close(); zin.close()
shutil.move(tmp, OUT)
print("deck saved:", OUT)
